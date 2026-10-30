"""EP-14 US-14.14: llegadas tarde y faltas contra el horario cargado (D18).

Riesgos que cubre:
- minutos tarde mal contados (segundos, turno cortado, la segunda entrada);
- una falta marcada hoy (todavia puede llegar), en un franco o antes del alta;
- un franco trabajado tomado como tarde o falta;
- dias sin horario contados fuera del alta o despues de la baja;
- la tolerancia del mes justa contada como pasada;
- que el Excel no muestre lo que dice el resumen.
"""

from datetime import date, datetime, time, timedelta
from io import BytesIO
from unittest import mock

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .excel import generar_excel, observacion_del_dia
from .models import Empleado, Horario, HorarioCambio, Marca
from .resumen import calcular_meses
from .services import dar_de_baja_empleado, guardar_reglas, reactivar_empleado
from .tests_excel import ExcelTestCase
from .tests_fichaje import _momento

HOY = date(2026, 11, 10)


class TardanzasYFaltasTestCase(ExcelTestCase):
    def setUp(self):
        super().setUp()
        self._alta(self.empleado, date(2026, 10, 1))

    def _alta(self, empleado, dia):
        Empleado.objects.filter(pk=empleado.pk).update(creado_en=timezone.make_aware(datetime.combine(dia, time(9))))
        empleado.refresh_from_db()

    def _horario(self, dia, *entradas, franco=False, empleado=None):
        entradas = (*entradas, None, None)
        return Horario.objects.create(
            empleado=empleado or self.empleado,
            fecha=date(2026, 11, dia),
            franco=franco,
            entrada_1=None if franco else entradas[0],
            entrada_2=None if franco else entradas[1],
        )

    def _resumen(self, hoy=HOY, empleado=None):
        empleado = empleado or self.empleado
        return calcular_meses([(empleado, list(empleado.marcas.all()))], 2026, 11, hoy=hoy)[0]

    def _dia(self, resumen, dia):
        return next(d for d in resumen.dias if d.fecha == date(2026, 11, dia))

    def _reloj(self, dia, hora, minuto=0):
        """Fija la hora de ahora en noviembre de 2026."""
        return mock.patch("django.utils.timezone.now", return_value=_momento(dia, hora, minuto))


class LlegadasTardeTests(TardanzasYFaltasTestCase):
    def test_a_horario_y_un_minuto_tarde(self):
        self._horario(2, time(8))
        self._turno(2, 8, 16)
        self._horario(3, time(8))
        self._marca(_momento(3, 8, 1), Marca.Tipo.ENTRADA)
        self._horario(4, time(8))
        self._marca(_momento(4, 7, 50), Marca.Tipo.ENTRADA)
        resumen = self._resumen()
        self.assertEqual(self._dia(resumen, 2).tarde, timedelta())
        self.assertEqual(self._dia(resumen, 3).tarde, timedelta(minutes=1))
        self.assertEqual(self._dia(resumen, 4).tarde, timedelta())
        self.assertEqual((resumen.llegadas_tarde, resumen.minutos_tarde), (1, 1))

    def test_los_segundos_no_cuentan(self):
        self._horario(2, time(8))
        self._marca(_momento(2, 8) + timedelta(seconds=59), Marca.Tipo.ENTRADA)
        self.assertEqual(self._dia(self._resumen(), 2).tarde, timedelta())

    def test_turno_cortado_compara_cada_entrada_con_la_suya(self):
        self._horario(2, time(8), time(17))
        self._marca(_momento(2, 8, 5), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 12), Marca.Tipo.SALIDA)
        self._marca(_momento(2, 17, 10), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 21), Marca.Tipo.SALIDA)
        dia = self._dia(self._resumen(), 2)
        self.assertEqual(dia.tarde, timedelta(minutes=15))
        self.assertEqual(dia.llegadas_tarde, 2)
        self.assertEqual(dia.turnos_sin_fichar, [])

    def test_vino_a_un_solo_turno_no_es_falta_y_se_senala(self):
        self._horario(2, time(8), time(17))
        self._turno(2, 8, 12)
        dia = self._dia(self._resumen(), 2)
        self.assertEqual(dia.tarde, timedelta())
        self.assertFalse(dia.falta)
        self.assertEqual(dia.turnos_sin_fichar, [2])
        self.assertEqual(observacion_del_dia(dia), "No fichó el 2.º turno")

    def test_vino_solo_al_segundo_turno_se_compara_con_el_segundo(self):
        # Antes se emparejaba por orden: 17:05 contra 8:00 daba 545 minutos tarde.
        self._horario(2, time(8), time(17))
        self._marca(_momento(2, 17, 5), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 21), Marca.Tipo.SALIDA)
        dia = self._dia(self._resumen(), 2)
        self.assertEqual(dia.tarde, timedelta(minutes=5))
        self.assertEqual(dia.turnos_sin_fichar, [1])
        self.assertFalse(dia.falta)
        self.assertEqual(observacion_del_dia(dia), "No fichó el 1.º turno")

    def test_una_reentrada_a_media_manana_no_tapa_la_tardanza_de_la_tarde(self):
        self._horario(2, time(8), time(17))
        for hora, minuto, tipo in ((8, 0, Marca.Tipo.ENTRADA), (10, 0, Marca.Tipo.SALIDA), (10, 20, Marca.Tipo.ENTRADA),
                                   (12, 0, Marca.Tipo.SALIDA), (17, 45, Marca.Tipo.ENTRADA), (21, 0, Marca.Tipo.SALIDA)):
            self._marca(_momento(2, hora, minuto), tipo)
        dia = self._dia(self._resumen(), 2)
        self.assertEqual(dia.tarde, timedelta(minutes=45))
        self.assertEqual(dia.llegadas_tarde, 1)

    def test_cada_entrada_tarde_es_una_llegada_tarde(self):
        self._horario(2, time(8), time(17))
        self._marca(_momento(2, 8, 10), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 12), Marca.Tipo.SALIDA)
        self._marca(_momento(2, 17, 10), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 21), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertEqual((resumen.llegadas_tarde, resumen.minutos_tarde), (2, 20))

    def test_un_turno_despues_de_medianoche_toma_al_que_llego_antes(self):
        # Entra a las 0:30 del 3; llego 23:55 del 2. No es falta ni llega tarde.
        self._horario(3, time(0, 30))
        self._marca(_momento(2, 23, 55), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 8), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        dia = self._dia(resumen, 3)
        self.assertEqual((dia.tarde, dia.falta), (timedelta(), False))
        self.assertEqual(dia.turnos[0].real, _momento(2, 23, 55))
        # Y si llega 0:40 del 3, son 10 minutos.
        self._horario(5, time(0, 30))
        self._marca(_momento(5, 0, 40), Marca.Tipo.ENTRADA)
        self.assertEqual(self._dia(self._resumen(), 5).tarde, timedelta(minutes=10))

    def test_turnos_de_noche_pegados_a_medianoche(self):
        # 23:00 del 2 y 0:00 del 3: la de 23:50 es del turno de las 0:00, la de 23:10 del de las 23:00.
        self._horario(2, time(23))
        self._horario(3, time(0))
        self._marca(_momento(2, 23, 10), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 23, 50), Marca.Tipo.ENTRADA)
        resumen = self._resumen()
        self.assertEqual(self._dia(resumen, 2).tarde, timedelta(minutes=10))
        self.assertEqual(self._dia(resumen, 3).tarde, timedelta())
        self.assertFalse(self._dia(resumen, 3).falta)

    def test_fichar_lejos_del_horario_no_es_tarde_ni_falta_y_se_senala(self):
        self._horario(2, time(8))
        self._marca(_momento(2, 21), Marca.Tipo.ENTRADA)
        dia = self._dia(self._resumen(), 2)
        self.assertEqual((dia.tarde, dia.falta), (timedelta(), False))
        self.assertEqual(observacion_del_dia(dia), "Fichó fuera de su horario, Sin salida")

    def test_una_entrada_de_la_noche_anterior_no_tapa_la_tardanza_ni_la_falta(self):
        # Segunda revision: con franjas de 12 h, la vuelta de las 21 del 2 contaba como llegada del 3.
        for dia in (2, 3, 4):
            self._horario(dia, time(8))
        self._turno(2, 8, 16)
        self._turno(2, 21, 23)
        self._marca(_momento(3, 8, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 16), Marca.Tipo.SALIDA)
        self._turno(3, 21, 23)
        resumen = self._resumen()
        self.assertEqual(self._dia(resumen, 3).tarde, timedelta(minutes=30))
        self.assertTrue(self._dia(resumen, 4).falta)

    def test_faltar_a_un_turno_de_noche_y_venir_temprano_al_otro_dia_es_falta(self):
        self._horario(2, time(20))
        self._horario(3, time(8))
        self._marca(_momento(3, 7), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 15), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        dia_2 = self._dia(resumen, 2)
        self.assertEqual((dia_2.falta, dia_2.tarde), (True, timedelta()))
        self.assertEqual(self._dia(resumen, 3).tarde, timedelta())

    def test_la_franja_de_un_turno_de_noche_no_pasa_de_las_5(self):
        # Tercera revision: con 8 h, el turno de las 22 tomaba la entrada de las 5:30 del otro dia.
        self._horario(2, time(22))
        self._horario(3, franco=True)
        self._marca(_momento(3, 5, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 13, 30), Marca.Tipo.SALIDA)
        # 22 del 5 y 14 del 6: trabajo el 6 de 5:30 a 13:30. El 5 es falta; el 6 no.
        self._horario(5, time(22))
        self._horario(6, time(14))
        self._marca(_momento(6, 5, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(6, 13, 30), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertEqual([self._dia(resumen, d).falta for d in (2, 5, 6)], [True, True, False])
        self.assertEqual((resumen.faltas, resumen.minutos_tarde), (2, 0))
        self.assertEqual(observacion_del_dia(self._dia(resumen, 6)), "Fichó fuera de su horario")
        # Hasta las 5 sigue siendo una llegada tarde: a las 0:30 son 150 minutos.
        self._horario(8, time(22))
        self._marca(_momento(9, 0, 30), Marca.Tipo.ENTRADA)
        self.assertEqual(self._dia(self._resumen(), 8).tarde, timedelta(minutes=150))

    def test_un_mes_cerrado_no_cambia_al_cargar_el_dia_1_del_siguiente(self):
        self._horario(30, time(22))
        self._marca(timezone.make_aware(datetime(2026, 12, 1, 5, 30)), Marca.Tipo.ENTRADA)
        antes = self._dia(self._resumen(hoy=date(2026, 12, 5)), 30)
        Horario.objects.create(empleado=self.empleado, fecha=date(2026, 12, 1), entrada_1=time(6))
        despues = self._dia(self._resumen(hoy=date(2026, 12, 5)), 30)
        for dia in (antes, despues):
            self.assertEqual((dia.falta, dia.tarde, observacion_del_dia(dia)), (True, timedelta(), "Falta"))

    def test_un_franco_trabajado_a_la_noche_no_tapa_la_falta_del_otro_dia(self):
        self._horario(2, franco=True)
        self._horario(3, time(7))
        self._marca(_momento(2, 19, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 23, 30), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertTrue(self._dia(resumen, 3).falta)
        self.assertEqual(resumen.faltas, 1)
        self.assertEqual(observacion_del_dia(self._dia(resumen, 2)), "Franco")

    def test_la_noche_del_ultimo_dia_del_mes_anterior_no_tapa_la_falta_del_1(self):
        # Por la vista, que trae las marcas de alrededor del mes.
        self._horario(1, time(7))
        self._marca(timezone.make_aware(datetime(2026, 10, 31, 19, 30)), Marca.Tipo.ENTRADA)
        self._marca(timezone.make_aware(datetime(2026, 10, 31, 23, 30)), Marca.Tipo.SALIDA)
        self._login(self.admin, [self.armadi])
        with self._reloj(10, 12):
            response = self.client.get(reverse("asistencia:excel"), {"descargar": "1", "mes": "2026-11"})
        libro = load_workbook(BytesIO(response.content))
        self.assertEqual(libro["Juan Perez"]["J6"].value, "Falta")
        fila = next(r for r in libro["Resumen"].iter_rows(min_row=6, values_only=True) if r[0] == "Juan Perez")
        self.assertEqual(fila[9], 1)

    def test_llegar_mas_de_dos_horas_antes_no_es_tarde_ni_se_senala(self):
        # 9 y 17, vuelve 14:30 y se queda hasta las 21. Horario 10, entra 7:45.
        self._horario(2, time(9), time(17))
        self._turno(2, 9, 13)
        self._marca(_momento(2, 14, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 21), Marca.Tipo.SALIDA)
        self._horario(3, time(10))
        self._marca(_momento(3, 7, 45), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 16), Marca.Tipo.SALIDA)
        # 9 y 17, pero se fue a las 16 sin quedarse al 2.º turno: ese si se senala.
        self._horario(4, time(9), time(17))
        self._turno(4, 9, 13)
        self._marca(_momento(4, 14, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(4, 16), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertEqual([observacion_del_dia(self._dia(resumen, d)) for d in (2, 3, 4)], ["", "", "No fichó el 2.º turno"])
        self.assertEqual((resumen.minutos_tarde, resumen.faltas), (0, 0))

    def test_panaderia_de_madrugada_no_toma_la_vuelta_de_la_tarde_anterior(self):
        for dia in (2, 3):
            self._horario(dia, time(4))
        self._turno(2, 4, 12)
        self._turno(2, 16, 18)
        self._marca(_momento(3, 4, 20), Marca.Tipo.ENTRADA)
        self.assertEqual(self._dia(self._resumen(), 3).tarde, timedelta(minutes=20))

    def test_volver_temprano_del_almuerzo_no_tapa_la_tardanza_de_la_tarde(self):
        self._horario(2, time(9), time(17))
        for hora, minuto, tipo in ((9, 0, Marca.Tipo.ENTRADA), (13, 0, Marca.Tipo.SALIDA), (13, 10, Marca.Tipo.ENTRADA),
                                   (14, 0, Marca.Tipo.SALIDA), (17, 30, Marca.Tipo.ENTRADA), (21, 0, Marca.Tipo.SALIDA)):
            self._marca(_momento(2, hora, minuto), tipo)
        self.assertEqual(self._dia(self._resumen(), 2).tarde, timedelta(minutes=30))

    def test_un_turno_de_madrugada_salteado_entre_dos_trabajados_es_falta(self):
        # 0:30 del 3, del 4 y del 5; vino el 3 (23:55 del 2) y el 5 (23:55 del 4), no el 4.
        for dia in (3, 4, 5):
            self._horario(dia, time(0, 30))
        self._marca(_momento(2, 23, 55), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 8), Marca.Tipo.SALIDA)
        self._marca(_momento(4, 23, 55), Marca.Tipo.ENTRADA)
        self._marca(_momento(5, 8), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertEqual([self._dia(resumen, d).falta for d in (3, 4, 5)], [False, True, False])
        self.assertEqual(resumen.faltas, 1)
        self.assertEqual(observacion_del_dia(self._dia(resumen, 3)), "Fichó la entrada el día anterior")

    def test_franco_trabajado_no_es_tarde_ni_falta(self):
        self._horario(2, franco=True)
        self._turno(2, 10, 14)
        dia = self._dia(self._resumen(), 2)
        self.assertEqual((dia.tarde, dia.falta), (timedelta(), False))
        self.assertEqual(observacion_del_dia(dia), "Franco")

    def test_la_tolerancia_justa_no_pasa_y_un_minuto_mas_si(self):
        guardar_reglas(actor=self.admin, empresa=self.armadi, jornada_minutos=510, tolerancia_mensual_minutos=30)
        for dia in (2, 3):
            self._horario(dia, time(8))
            self._marca(_momento(dia, 8, 15), Marca.Tipo.ENTRADA)
        resumen = self._resumen()
        self.assertEqual(resumen.minutos_tarde, 30)
        self.assertFalse(resumen.pasa_la_tolerancia)
        self._horario(4, time(8))
        self._marca(_momento(4, 8, 1), Marca.Tipo.ENTRADA)
        resumen = self._resumen()
        self.assertEqual(resumen.minutos_tarde, 31)
        self.assertTrue(resumen.pasa_la_tolerancia)


class EntradasDeOtroDiaTests(TardanzasYFaltasTestCase):
    """Una entrada que le toca al turno de otro dia queda en la fila del dia en que
    se ficho: la hoja lo dice en las dos filas (tercera revision de D18)."""

    def test_turnos_de_madrugada_seguidos_lo_dicen_en_las_dos_filas(self):
        for dia in (3, 4, 5):
            self._horario(dia, time(0, 30))
        for dia in (2, 3, 4):
            self._marca(_momento(dia, 23, 55), Marca.Tipo.ENTRADA)
            self._marca(_momento(dia + 1, 8), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        siguiente = "La entrada de las 23:55 es del turno del día siguiente"
        self.assertEqual(
            [observacion_del_dia(self._dia(resumen, d)) for d in (2, 3, 4, 5)],
            [
                f"Sin horario, {siguiente}",
                f"Fichó la entrada el día anterior, {siguiente}",
                f"Fichó la entrada el día anterior, {siguiente}",
                "Fichó la entrada el día anterior",
            ],
        )
        self.assertEqual((resumen.faltas, resumen.minutos_tarde), (0, 0))

    def test_una_falta_que_muestra_la_entrada_del_turno_del_dia_anterior_lo_dice(self):
        # 22 del 6 y del 7: el 6 llego a las 0:30 del 7 (150 minutos tarde) y el 7 no vino.
        self._horario(6, time(22))
        self._horario(7, time(22))
        self._marca(_momento(7, 0, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(7, 6), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertEqual(observacion_del_dia(self._dia(resumen, 6)), "Fichó la entrada el día siguiente")
        self.assertEqual(
            observacion_del_dia(self._dia(resumen, 7)), "Falta, La entrada de las 00:30 es del turno del día anterior"
        )
        self.assertEqual((resumen.faltas, resumen.minutos_tarde), (1, 150))

    def test_una_salida_suelta_con_la_entrada_de_otro_dia_no_dice_fuera_de_horario(self):
        # El 9: la entrada de las 8 se anulo y quedo la salida; a las 23:30 entro al turno de las 0 del 10.
        self._horario(9, time(8))
        self._horario(10, time(0))
        self._marca(_momento(9, 8), Marca.Tipo.ENTRADA, anulada=True)
        self._marca(_momento(9, 16), Marca.Tipo.SALIDA)
        self._marca(_momento(9, 23, 30), Marca.Tipo.ENTRADA)
        self._marca(_momento(10, 7), Marca.Tipo.SALIDA)
        dia = self._dia(self._resumen(hoy=date(2026, 11, 12)), 9)
        self.assertEqual((dia.turnos_sin_fichar, dia.falta), ([], False))
        self.assertEqual(
            observacion_del_dia(dia), "La entrada de las 23:30 es del turno del día siguiente, Salida sin entrada"
        )


class FaltasYSinHorarioTests(TardanzasYFaltasTestCase):
    def test_falta_solo_hasta_ayer(self):
        for dia in (8, 9, 10, 11):
            self._horario(dia, time(8))
        resumen = self._resumen(hoy=date(2026, 11, 10))
        self.assertEqual([self._dia(resumen, d).falta for d in (8, 9, 10, 11)], [True, True, False, False])
        self.assertEqual(resumen.faltas, 2)
        self.assertEqual(observacion_del_dia(self._dia(resumen, 9)), "Falta")

    def test_hoy_no_cuenta_como_sin_salida_ni_turno_sin_fichar(self):
        self._horario(10, time(8), time(17))
        self._marca(_momento(10, 8), Marca.Tipo.ENTRADA)
        hoy = self._resumen(hoy=date(2026, 11, 10))
        self.assertEqual((hoy.dias_sin_salida, self._dia(hoy, 10).turnos_sin_fichar), (0, []))
        manana = self._resumen(hoy=date(2026, 11, 11))
        self.assertEqual((manana.dias_sin_salida, self._dia(manana, 10).turnos_sin_fichar), (1, [2]))

    def test_hoy_la_hoja_dice_en_curso_y_no_sin_salida(self):
        self._horario(10, time(8))
        self._marca(_momento(10, 8), Marca.Tipo.ENTRADA)
        self.assertEqual(observacion_del_dia(self._dia(self._resumen(hoy=date(2026, 11, 10)), 10)), "En curso")
        self.assertEqual(observacion_del_dia(self._dia(self._resumen(hoy=date(2026, 11, 11)), 10)), "Sin salida")

    def test_una_salida_suelta_no_dice_fuera_de_horario(self):
        self._horario(2, time(8))
        self._marca(_momento(2, 16), Marca.Tipo.SALIDA)
        self.assertEqual(observacion_del_dia(self._dia(self._resumen(), 2)), "Salida sin entrada")

    def test_una_salida_sin_entrada_no_es_falta(self):
        self._horario(2, time(8))
        self._marca(_momento(2, 16), Marca.Tipo.SALIDA)
        self.assertFalse(self._dia(self._resumen(), 2).falta)

    def test_un_franco_sin_marcas_no_es_falta(self):
        self._horario(2, franco=True)
        dia = self._dia(self._resumen(), 2)
        self.assertFalse(dia.falta)
        self.assertEqual(observacion_del_dia(dia), "Franco")

    def test_dias_sin_horario_desde_el_alta_hasta_ayer(self):
        self._alta(self.empleado, date(2026, 11, 5))
        self._horario(6, time(8))
        self._turno(6, 8, 16)
        resumen = self._resumen(hoy=date(2026, 11, 10))
        # 5, 7, 8 y 9 sin horario; el 6 tiene; antes del 5 no estaba; el 10 es hoy.
        self.assertEqual(resumen.dias_sin_horario, 4)
        self.assertFalse(self._dia(resumen, 4).sin_horario)
        self.assertEqual(observacion_del_dia(self._dia(resumen, 4)), "Sin marcas")
        self.assertEqual(observacion_del_dia(self._dia(resumen, 7)), "Sin horario")
        # Un horario cargado antes del alta tampoco es falta.
        self._horario(3, time(8))
        self.assertFalse(self._dia(self._resumen(hoy=date(2026, 11, 10)), 3).falta)

    def test_despues_de_la_baja_no_cuenta(self):
        self._horario(9, time(8))
        dar_de_baja_empleado(actor=self.admin, empleado=self.empleado)
        Empleado.objects.filter(pk=self.empleado.pk).update(
            dado_de_baja_en=timezone.make_aware(datetime(2026, 11, 7, 12))
        )
        self.empleado.refresh_from_db()
        resumen = self._resumen(hoy=date(2026, 11, 20))
        self.assertEqual(resumen.dias_sin_horario, 7)
        self.assertFalse(self._dia(resumen, 9).falta)


class BajasYReactivacionesTests(TardanzasYFaltasTestCase):
    def test_la_baja_no_borra_horarios_y_al_volver_se_limpian_los_dias_de_baja(self):
        fechas = [date(2026, 11, dia) for dia in (2, 7, 10, 12)]
        for fecha in fechas:
            Horario.objects.create(empleado=self.empleado, fecha=fecha, entrada_1=time(8))
        with self._reloj(5, 12):
            dar_de_baja_empleado(actor=self.admin, empleado=self.empleado)
        # La baja no toca la grilla: una baja por error se deshace sin perder nada cargado.
        self.assertEqual(sorted(Horario.objects.values_list("fecha", flat=True)), fechas)
        self.assertFalse(HorarioCambio.objects.exists())
        # Vuelve el 10 a las 7: el horario del 7 (estaba de baja) no le hace una falta; el de hoy a las 8 queda.
        with self._reloj(10, 7):
            reactivar_empleado(actor=self.admin, empleado=self.empleado)
        self.assertEqual(sorted(Horario.objects.values_list("fecha", flat=True)), [fechas[0], fechas[2], fechas[3]])
        self.assertEqual(
            list(HorarioCambio.objects.values_list("fecha", "antes", "despues", "cambiado_por")),
            [(fechas[1], "8:00", "", self.admin.pk)],
        )

    def test_al_volver_a_media_manana_no_cuentan_los_turnos_de_hoy_que_ya_empezaron(self):
        # Se fueron el 2 y vuelven el 5 a las 12: Juan tenia 8 y 17, Ana 8 y Beto franco.
        ana = self._empleado(nombre="Ana Gomez", dni="30999888")
        beto = self._empleado(nombre="Beto Franco", dni="31000111")
        self._alta(ana, date(2026, 10, 1))
        self._horario(5, time(8), time(17))
        self._horario(5, time(8), empleado=ana)
        self._horario(5, franco=True, empleado=beto)
        for empleado in (self.empleado, ana, beto):
            with self._reloj(2, 9, 30):
                dar_de_baja_empleado(actor=self.admin, empleado=empleado)
            with self._reloj(5, 12):
                reactivar_empleado(actor=self.admin, empleado=empleado)
        self.assertEqual(Horario.objects.get(empleado=self.empleado).entradas, (time(17),))
        self.assertFalse(Horario.objects.filter(empleado=ana).exists())
        self.assertTrue(Horario.objects.get(empleado=beto).franco)
        self.assertEqual(
            sorted(HorarioCambio.objects.values_list("empleado__nombre", "antes", "despues")),
            [("Ana Gomez", "8:00", ""), ("Juan Perez", "8:00 y 17:00", "17:00")],
        )
        # El 5 no es falta para Ana: a las 8 no trabajaba aca.
        self.assertFalse(self._dia(self._resumen(hoy=date(2026, 11, 6), empleado=ana), 5).falta)

    def test_una_baja_y_vuelta_el_mismo_dia_no_toca_el_horario_de_hoy(self):
        self._horario(5, time(8))
        with self._reloj(5, 9, 30):
            dar_de_baja_empleado(actor=self.admin, empleado=self.empleado)
        with self._reloj(5, 9, 35):
            reactivar_empleado(actor=self.admin, empleado=self.empleado)
        self.assertEqual(Horario.objects.get(empleado=self.empleado).entradas, (time(8),))
        self.assertFalse(HorarioCambio.objects.exists())

    def test_una_baja_y_vuelta_inmediata_no_pierde_la_grilla(self):
        hoy = timezone.localdate()
        for dias in (1, 2, 3):
            Horario.objects.create(empleado=self.empleado, fecha=hoy + timedelta(days=dias), entrada_1=time(8))
        dar_de_baja_empleado(actor=self.admin, empleado=self.empleado)
        reactivar_empleado(actor=self.admin, empleado=self.empleado)
        self.assertEqual(Horario.objects.count(), 3)

    def test_un_dado_de_baja_sin_marcas_pero_con_faltas_sale_en_el_excel(self):
        # Octubre ya paso. Ana tenia horario el 6 y el 7, no vino y se fue el 20.
        ana = self._empleado(nombre="Ana Baja", dni="32000111")
        sin_faltas = self._empleado(nombre="Beto Franco", dni="33000111")
        for empleado in (ana, sin_faltas):
            self._alta(empleado, date(2026, 9, 1))
        Horario.objects.create(empleado=ana, fecha=date(2026, 10, 6), entrada_1=time(8))
        Horario.objects.create(empleado=ana, fecha=date(2026, 10, 7), entrada_1=time(8))
        Horario.objects.create(empleado=sin_faltas, fecha=date(2026, 10, 6), franco=True)
        for empleado in (ana, sin_faltas):
            dar_de_baja_empleado(actor=self.admin, empleado=empleado)
            Empleado.objects.filter(pk=empleado.pk).update(dado_de_baja_en=timezone.make_aware(datetime(2026, 10, 20, 12)))
        self._login(self.admin, [self.armadi])
        response = self.client.get(reverse("asistencia:excel"), {"descargar": "1", "mes": "2026-10"})
        libro = load_workbook(BytesIO(response.content))
        self.assertEqual(libro.sheetnames, ["Resumen", "Ana Baja", "Juan Perez"])
        fila = next(r for r in libro["Resumen"].iter_rows(min_row=6, values_only=True) if r[0] == "Ana Baja")
        self.assertEqual(fila[9], 2)


class ExcelConTardanzasTests(TardanzasYFaltasTestCase):
    def test_la_hoja_y_el_resumen_muestran_horario_tarde_y_faltas(self):
        guardar_reglas(actor=self.admin, empresa=self.armadi, jornada_minutos=510, tolerancia_mensual_minutos=10)
        self._horario(2, time(8))
        self._marca(_momento(2, 8, 20), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 16), Marca.Tipo.SALIDA)
        self._horario(3, time(8), time(17))
        self._horario(4, franco=True)
        libro = load_workbook(BytesIO(generar_excel([(self.empleado, list(self.empleado.marcas.all()))], 2026, 11, hoy=date(2026, 11, 5))))
        hoja = libro["Juan Perez"]
        encabezados = [hoja.cell(row=5, column=c).value for c in range(1, 11)]
        self.assertEqual(
            encabezados,
            ["Día", "Horario", "Entrada 1", "Salida 1", "Entrada 2", "Salida 2", "Horas del día", "Horas extra", "Tarde (min)", "Observación"],
        )
        dia_2, dia_3, dia_4 = (hoja[f"{columna}{fila}"] for columna, fila in (("B", 7), ("B", 8), ("B", 9)))
        self.assertEqual((dia_2.value, dia_3.value, dia_4.value), ("8:00", "8:00 y 17:00", "Franco"))
        self.assertEqual(hoja["I7"].value, 20)
        self.assertEqual(hoja["J8"].value, "Falta")
        self.assertEqual(hoja["J9"].value, "Franco")
        self.assertEqual(hoja["J6"].value, "Sin horario")
        self.assertEqual(hoja["I36"].value, 20)

        resumen = libro["Resumen"]
        fila = [resumen.cell(row=6, column=c).value for c in range(7, 12)]
        # Llegadas tarde, minutos, pasa la tolerancia, faltas, dias sin horario (el 1).
        self.assertEqual(fila, [1, 20, "Sí", 1, 1])
        self.assertIn("ARMADI 10 min", resumen["A3"].value)
