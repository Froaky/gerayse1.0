"""Excel del mes por empleado (EP-14 US-14.6).

Una hoja por empleado con todos los dias del mes: entradas y salidas en pares,
horas del dia y una observacion. Las reglas de tramos estan en `jornadas.py` y
las del resumen (horas extra, D18) en `resumen.py`: aca solo se dibuja. Una
marca cargada a mano sale como "Corregida".
"""

import re
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

# Tramo, armar_tramos, jornadas_del_mes, horas_del_dia y ventana_del_mes se
# siguen importando desde aca (vistas, Por local y tests).
from .jornadas import Tramo, armar_tramos, horas_del_dia, jornadas_del_mes, ventana_del_mes  # noqa: F401
from .resumen import ResumenMes, calcular_meses

MESES = [
    "", "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
]
CONTENT_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def observacion(tramos: list[Tramo]) -> str:
    if not tramos:
        return "Sin marcas"
    notas = []
    if any(t.corregido for t in tramos):
        notas.append("Corregida")
    if any(t.entrada is not None and t.salida is None for t in tramos):
        notas.append("Sin salida")
    if any(t.entrada is None for t in tramos):
        notas.append("Salida sin entrada")
    return ", ".join(notas)


def _recortar(texto: str, maximo: int) -> str:
    # Excel cuenta el largo en unidades UTF-16: un emoji ocupa dos. Tampoco
    # acepta un apostrofo al principio o al final.
    while len(texto.encode("utf-16-le")) // 2 > maximo:
        texto = texto[:-1]
    return texto.strip(" '")


def _titulo_hoja(nombre: str, usados: set) -> str:
    # Excel: hasta 31 caracteres, sin []:*?/\ y sin repetir. Un nombre de hoja
    # invalido hace que Excel no abra el libro entero.
    base = _recortar(re.sub(r"[\[\]:*?/\\]", " ", ILLEGAL_CHARACTERS_RE.sub("", nombre)).strip(" '"), 31) or "Empleado"
    titulo, numero = base, 2
    while titulo.lower() in usados:
        sufijo = f" ({numero})"
        titulo = (_recortar(base, 31 - len(sufijo)) or "Empleado") + sufijo
        numero += 1
    usados.add(titulo.lower())
    return titulo


def _texto(hoja, coordenada: str, valor: str):
    # Un nombre que empieza con "=" no tiene que volverse una formula.
    celda = hoja[coordenada]
    celda.value = ILLEGAL_CHARACTERS_RE.sub("", valor)
    celda.data_type = "s"
    return celda


def _hora(marca) -> object:
    return timezone.localtime(marca.momento).time().replace(second=0, microsecond=0) if marca else None


ENCABEZADO = PatternFill("solid", fgColor="1B3A2E")
ENLACE = Font(color="0563C1", underline="single")
TITULO_RESUMEN = "Resumen"


def _encabezar(hoja, fila: int, titulos: list[str]) -> None:
    for columna, titulo in enumerate(titulos, start=1):
        celda = hoja.cell(row=fila, column=columna, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _duracion(hoja, fila: int, columna: int, valor, negrita: bool = False):
    celda = hoja.cell(row=fila, column=columna, value=valor)
    celda.number_format = "[h]:mm"
    celda.alignment = Alignment(horizontal="center")
    if negrita:
        celda.font = Font(bold=True)
    return celda


def _local(empleado) -> str:
    return empleado.sucursal.nombre if empleado.sucursal.empresa_id == empleado.empresa_id else "local de otra empresa"


def _hoja_empleado(hoja, resumen: ResumenMes, anio: int, mes: int) -> None:
    empleado = resumen.empleado
    pares = max([2, *(len(dia.tramos) for dia in resumen.dias)])
    _texto(hoja, "A1", empleado.nombre).font = Font(bold=True, size=14)
    _texto(hoja, "A2", f"{empleado.empresa.nombre} · {_local(empleado)}" + ("" if empleado.activo else " · dado de baja"))
    hoja["A3"] = f"Asistencia de {MESES[mes]} de {anio}"
    hoja["A4"] = f"Horas extra: lo que pasa de {resumen.reglas.jornada_texto} por día."

    titulos = ["Día"]
    for numero in range(1, pares + 1):
        titulos += [f"Entrada {numero}", f"Salida {numero}"]
    titulos += ["Horas del día", "Horas extra", "Observación"]
    _encabezar(hoja, 5, titulos)
    columna_horas = 2 + pares * 2

    fila = 6
    for dia in resumen.dias:
        hoja.cell(row=fila, column=1, value=dia.fecha).number_format = "ddd dd/mm"
        for indice, tramo in enumerate(dia.tramos):
            for desplazamiento, marca in ((0, tramo.entrada), (1, tramo.salida)):
                celda = hoja.cell(row=fila, column=2 + indice * 2 + desplazamiento, value=_hora(marca))
                celda.number_format = "hh:mm"
                celda.alignment = Alignment(horizontal="center")
        _duracion(hoja, fila, columna_horas, dia.horas if dia.tramos else None)
        _duracion(hoja, fila, columna_horas + 1, dia.extra or None)
        hoja.cell(row=fila, column=columna_horas + 2, value=observacion(dia.tramos))
        fila += 1
    hoja.cell(row=fila, column=1, value="Total del mes").font = Font(bold=True)
    _duracion(hoja, fila, columna_horas, resumen.horas, negrita=True)
    _duracion(hoja, fila, columna_horas + 1, resumen.extra, negrita=True)

    hoja.column_dimensions["A"].width = 14
    for columna in range(2, columna_horas):
        hoja.column_dimensions[get_column_letter(columna)].width = 11
    hoja.column_dimensions[get_column_letter(columna_horas)].width = 14
    hoja.column_dimensions[get_column_letter(columna_horas + 1)].width = 13
    hoja.column_dimensions[get_column_letter(columna_horas + 2)].width = 28
    hoja.freeze_panes = "B6"


def _hoja_resumen(hoja, resumenes: list[ResumenMes], titulos: dict, anio: int, mes: int) -> None:
    """Una fila por empleado, por local y nombre, con filtro: el encargado filtra
    su local. El nombre lleva a la hoja de esa persona."""
    hoja["A1"] = f"Resumen de asistencia de {MESES[mes]} de {anio}"
    hoja["A1"].font = Font(bold=True, size=14)
    jornadas = sorted({(r.empleado.empresa.nombre, r.reglas.jornada_texto) for r in resumenes})
    _texto(hoja, "A2", "Horas extra: lo que pasa de la jornada diaria de cada empresa (" + ", ".join(f"{nombre} {jornada}" for nombre, jornada in jornadas) + ").")
    hoja["A3"] = "Hacé clic en un nombre para ver sus días. Con las flechas del encabezado filtrás por local."

    columnas = ["Empleado", "Empresa", "Local", "Días trabajados", "Horas del mes", "Horas extra", "Días sin salida", "Marcas corregidas"]
    _encabezar(hoja, 5, columnas)
    hoja.row_dimensions[5].height = 30
    orden = sorted(resumenes, key=lambda r: (_local(r.empleado).lower(), r.empleado.nombre.lower(), r.empleado.pk or 0))
    fila = 6
    for resumen in orden:
        empleado = resumen.empleado
        nombre = _texto(hoja, f"A{fila}", empleado.nombre)
        destino = titulos[id(resumen)].replace("'", "''")
        nombre.hyperlink = Hyperlink(ref=nombre.coordinate, location=f"'{destino}'!A1")
        nombre.font = ENLACE
        _texto(hoja, f"B{fila}", empleado.empresa.nombre)
        _texto(hoja, f"C{fila}", _local(empleado) + ("" if empleado.activo else " (de baja)"))
        hoja.cell(row=fila, column=4, value=resumen.dias_trabajados)
        _duracion(hoja, fila, 5, resumen.horas)
        _duracion(hoja, fila, 6, resumen.extra)
        hoja.cell(row=fila, column=7, value=resumen.dias_sin_salida)
        hoja.cell(row=fila, column=8, value=resumen.marcas_corregidas)
        for columna in (4, 7, 8):
            hoja.cell(row=fila, column=columna).alignment = Alignment(horizontal="center")
        fila += 1
    hoja.auto_filter.ref = f"A5:{get_column_letter(len(columnas))}{max(fila - 1, 5)}"
    for letra, ancho in zip("ABCDEFGH", (30, 16, 26, 11, 11, 11, 11, 11)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "B6"


def generar_excel(empleados_y_marcas, anio: int, mes: int) -> bytes:
    """`empleados_y_marcas`: [(empleado, marcas del mes y alrededores)]. La
    primera hoja es el Resumen (US-14.11); despues una por empleado."""
    resumenes = calcular_meses(empleados_y_marcas, anio, mes)
    libro = Workbook()
    libro.remove(libro.active)
    # "Resumen" se reserva primero: un empleado que se llame asi sale "Resumen (2)".
    usados = {TITULO_RESUMEN.lower()}
    titulos = {}
    for resumen in resumenes:
        hoja = libro.create_sheet(_titulo_hoja(resumen.empleado.nombre, usados))
        titulos[id(resumen)] = hoja.title
        _hoja_empleado(hoja, resumen, anio, mes)
    if resumenes:
        _hoja_resumen(libro.create_sheet(TITULO_RESUMEN, 0), resumenes, titulos, anio, mes)
        libro.active = 0
    else:
        libro.create_sheet("Sin empleados")["A1"] = "No hay empleados para ese mes."
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
