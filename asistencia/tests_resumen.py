"""EP-14 US-14.11: hoja Resumen y horas extra del Excel (D18).

Riesgos que cubre:
- horas extra mal contadas: un dia justo en la jornada, un minuto de mas, turno cortado;
- que un tramo sin salida sume horas o extra;
- que se use la jornada de otra empresa;
- que el Resumen no coincida con las hojas o salga desordenado;
- que un empleado llamado "Resumen" pise la hoja del resumen o el enlace lleve a otra hoja.
"""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from .excel import generar_excel
from .models import Marca
from .resumen import calcular_meses
from .services import guardar_reglas
from .tests_excel import ExcelTestCase
from .tests_fichaje import _momento

# Hoja de empleado: Dia, Horario y dos pares de entrada y salida.
COLUMNA_HORAS, COLUMNA_EXTRA, COLUMNA_TARDE, COLUMNA_OBSERVACION = 7, 8, 9, 10
FILA_DIA_1 = 6
# Con el mes ya cerrado: "sin salida" no cuenta el dia de hoy, que sigue en curso.
FIN_DE_MES = date(2026, 12, 1)


class HorasExtraTests(ExcelTestCase):
    def _resumen(self, empleado=None):
        empleado = empleado or self.empleado
        return calcular_meses([(empleado, list(empleado.marcas.all()))], 2026, 11, hoy=FIN_DE_MES)[0]

    def _dia(self, resumen, dia):
        return next(d for d in resumen.dias if d.fecha == _momento(dia, 0).date())

    def test_justo_la_jornada_no_es_extra_y_un_minuto_mas_si(self):
        self._marca(_momento(2, 8), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 16, 30), Marca.Tipo.SALIDA)
        self._marca(_momento(3, 8), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 16, 31), Marca.Tipo.SALIDA)
        resumen = self._resumen()
        self.assertEqual(self._dia(resumen, 2).extra, timedelta())
        self.assertEqual(self._dia(resumen, 3).extra, timedelta(minutes=1))
        self.assertEqual(resumen.extra, timedelta(minutes=1))
        self.assertEqual(resumen.horas, timedelta(hours=17, minutes=1))
        self.assertEqual(resumen.dias_trabajados, 2)

    def test_turno_cortado_suma_los_dos_tramos_antes_de_comparar(self):
        self._turno(4, 8, 13)
        self._turno(4, 17, 22)
        dia = self._dia(self._resumen(), 4)
        self.assertEqual(dia.horas, timedelta(hours=10))
        self.assertEqual(dia.extra, timedelta(hours=1, minutes=30))

    def test_un_tramo_sin_salida_no_suma_horas_ni_extra(self):
        self._marca(_momento(5, 6), Marca.Tipo.ENTRADA)
        resumen = self._resumen()
        dia = self._dia(resumen, 5)
        self.assertEqual((dia.horas, dia.extra), (timedelta(), timedelta()))
        self.assertTrue(dia.trabajado)
        self.assertEqual(resumen.dias_sin_salida, 1)

    def test_cada_empresa_con_su_jornada(self):
        guardar_reglas(actor=self.admin, empresa=self.mapogo, jornada_minutos=9 * 60, tolerancia_mensual_minutos=30)
        de_mapogo = self._empleado(nombre="Ana Vivre", dni="31000111", sucursal=self.vivre)
        self._turno(2, 8, 18)
        self._turno(2, 8, 18, empleado=de_mapogo)
        resumenes = calcular_meses(
            [(e, list(e.marcas.all())) for e in (self.empleado, de_mapogo)], 2026, 11
        )
        self.assertEqual([r.extra for r in resumenes], [timedelta(hours=1, minutes=30), timedelta(hours=1)])

    def test_cuenta_las_marcas_corregidas(self):
        self._marca(_momento(2, 8), Marca.Tipo.ENTRADA, origen=Marca.Origen.MANUAL)
        self._marca(_momento(2, 16), Marca.Tipo.SALIDA)
        self._marca(_momento(3, 8), Marca.Tipo.ENTRADA)
        self._marca(_momento(3, 16), Marca.Tipo.SALIDA, origen=Marca.Origen.MANUAL)
        self._marca(_momento(4, 8), Marca.Tipo.ENTRADA, origen=Marca.Origen.MANUAL, anulada=True)
        self.assertEqual(self._resumen().marcas_corregidas, 2)


class HojaResumenTests(ExcelTestCase):
    def _libro(self, empleados):
        contenido = generar_excel([(e, list(e.marcas.all())) for e in empleados], 2026, 11, hoy=FIN_DE_MES)
        return load_workbook(BytesIO(contenido))

    def test_la_hoja_del_empleado_trae_horas_extra_por_dia_y_total(self):
        self._turno(2, 8, 18)
        hoja = self._libro([self.empleado])["Juan Perez"]
        self.assertEqual(hoja.cell(row=5, column=COLUMNA_EXTRA).value, "Horas extra")
        self.assertEqual(hoja.cell(row=5, column=COLUMNA_OBSERVACION).value, "Observación")
        self.assertEqual(hoja.cell(row=FILA_DIA_1 + 1, column=COLUMNA_EXTRA).value, timedelta(hours=1, minutes=30))
        self.assertIsNone(hoja.cell(row=FILA_DIA_1, column=COLUMNA_EXTRA).value)
        self.assertEqual(hoja.cell(row=FILA_DIA_1 + 30, column=COLUMNA_EXTRA).value, timedelta(hours=1, minutes=30))
        self.assertTrue(hoja["A4"].value.startswith("Horas extra: lo que pasa de 8:30 por día."))

    def test_el_resumen_va_primero_por_local_y_nombre_y_coincide_con_las_hojas(self):
        de_vivre = self._empleado(nombre="Zoe Vivre", dni="31000111", sucursal=self.vivre)
        de_terminal = self._empleado(nombre="Beto Terminal", dni="32000111", sucursal=self.terminal)
        otro_centro = self._empleado(nombre="Ana Centro", dni="33000111")
        self._turno(2, 8, 18)
        self._turno(3, 8, 12)
        self._marca(_momento(4, 8), Marca.Tipo.ENTRADA)
        self._turno(2, 7, 15, empleado=de_vivre)
        libro = self._libro([otro_centro, de_terminal, self.empleado, de_vivre])
        self.assertEqual(libro.sheetnames[0], "Resumen")
        resumen = libro["Resumen"]
        self.assertEqual(resumen["A1"].value, "Resumen de asistencia de noviembre de 2026")
        encabezados = [resumen.cell(row=5, column=c).value for c in range(1, 14)]
        self.assertEqual(
            encabezados,
            [
                "Empleado", "Empresa", "Local", "Días trabajados", "Horas del mes", "Horas extra",
                "Llegadas tarde", "Minutos tarde", "Pasa la tolerancia", "Faltas", "Días sin horario",
                "Días sin salida", "Marcas corregidas",
            ],
        )
        filas = [[resumen.cell(row=f, column=c).value for c in range(1, 14)] for f in range(6, 10)]
        # Centro, Terminal, Vivré; dentro del local, por nombre.
        self.assertEqual([f[0] for f in filas], ["Ana Centro", "Juan Perez", "Beto Terminal", "Zoe Vivre"])
        juan = filas[1]
        self.assertEqual(juan[1:4], ["ARMADI", "Centro", 3])
        self.assertEqual(juan[4:6], [timedelta(hours=14), timedelta(hours=1, minutes=30)])
        self.assertEqual(juan[11:13], [1, 0])
        self.assertEqual(resumen.auto_filter.ref, "A5:M9")
        # El total de horas del Resumen es el mismo que el de la hoja de la persona.
        self.assertEqual(libro["Juan Perez"].cell(row=FILA_DIA_1 + 30, column=COLUMNA_HORAS).value, juan[4])

    def test_el_nombre_lleva_a_su_hoja_aunque_alguien_se_llame_resumen(self):
        resumen_persona = self._empleado(nombre="Resumen", dni="34000111")
        con_apostrofo = self._empleado(nombre="Juan D'Angelo", dni="35000111")
        libro = self._libro([self.empleado, resumen_persona, con_apostrofo])
        self.assertEqual(libro.sheetnames, ["Resumen", "Juan Perez", "Resumen (2)", "Juan D'Angelo"])
        hoja = libro["Resumen"]
        enlaces = {hoja.cell(row=f, column=1).value: hoja.cell(row=f, column=1).hyperlink.location for f in range(6, 9)}
        self.assertEqual(
            enlaces,
            {"Juan D'Angelo": "'Juan D''Angelo'!A1", "Juan Perez": "'Juan Perez'!A1", "Resumen": "'Resumen (2)'!A1"},
        )

    def test_dice_la_jornada_de_cada_empresa(self):
        guardar_reglas(actor=self.admin, empresa=self.mapogo, jornada_minutos=9 * 60, tolerancia_mensual_minutos=30)
        de_vivre = self._empleado(nombre="Zoe Vivre", dni="31000111", sucursal=self.vivre)
        libro = self._libro([self.empleado, de_vivre])
        self.assertEqual(
            libro["Resumen"]["A2"].value,
            "Horas extra: lo que pasa de la jornada diaria de cada empresa (ARMADI 8:30, MAPOGO 9:00).",
        )
        self.assertTrue(libro["Zoe Vivre"]["A4"].value.startswith("Horas extra: lo que pasa de 9:00 por día."))
