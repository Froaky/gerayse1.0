"""EP-14 US-14.6: Excel del mes por empleado.

Riesgos que cubre:
- horas mal sumadas en un turno cortado o una salida despues de medianoche;
- que una marca anulada cuente o que una corregida no se vea;
- que la salida del dia 1 que cierra una entrada del mes anterior aparezca;
- que entren empleados de otra empresa o falten los dados de baja con marcas;
- que un nombre que empieza con "=" se vuelva una formula.
"""

from datetime import timedelta
from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .excel import _titulo_hoja, generar_excel, jornadas_del_mes, observacion
from .models import Marca
from .services import dar_de_baja_empleado
from .tests_fichaje import FichajeTestCase, _momento

FILA_DIA_1 = 6


class ExcelTestCase(FichajeTestCase):
    def _marca(self, cuando, tipo, empleado=None, origen=Marca.Origen.ESCANEO, anulada=False):
        return Marca.objects.create(
            empleado=empleado or self.empleado,
            tipo=tipo,
            momento=cuando,
            sucursal=self.terminal,
            origen=origen,
            registrada_por=self.cajero,
            motivo="Se olvidó el carnet" if origen == Marca.Origen.MANUAL else "",
            anulada=anulada,
        )

    def _turno(self, dia, desde, hasta, empleado=None):
        self._marca(_momento(dia, desde), Marca.Tipo.ENTRADA, empleado)
        self._marca(_momento(dia, hasta), Marca.Tipo.SALIDA, empleado)

    def _jornadas(self, anio=2026, mes=11):
        return jornadas_del_mes(list(self.empleado.marcas.all()), anio, mes)


class ReglasDelExcelTests(ExcelTestCase):
    def test_turno_cortado_suma_los_dos_tramos(self):
        self._turno(4, 9, 13)
        self._turno(4, 18, 23)
        tramos = self._jornadas()[_momento(4, 0).date()]
        self.assertEqual(len(tramos), 2)
        self.assertEqual(sum((t.horas for t in tramos), timedelta()), timedelta(hours=9))
        self.assertEqual(observacion(tramos), "")

    def test_la_salida_de_madrugada_cuenta_en_el_dia_de_la_entrada(self):
        self._marca(_momento(5, 20), Marca.Tipo.ENTRADA)
        self._marca(_momento(6, 2), Marca.Tipo.SALIDA)
        jornadas = self._jornadas()
        self.assertEqual(jornadas[_momento(5, 0).date()][0].horas, timedelta(hours=6))
        self.assertEqual(jornadas[_momento(6, 0).date()], [])

    def test_sin_salida_no_suma_y_queda_senalado(self):
        self._marca(_momento(2, 8), Marca.Tipo.ENTRADA)
        tramos = self._jornadas()[_momento(2, 0).date()]
        self.assertIsNone(tramos[0].horas)
        self.assertEqual(observacion(tramos), "Sin salida")

    def test_una_marca_cargada_a_mano_sale_corregida(self):
        self._marca(_momento(2, 6), Marca.Tipo.ENTRADA, origen=Marca.Origen.MANUAL)
        self._marca(_momento(2, 14), Marca.Tipo.SALIDA)
        self.assertEqual(observacion(self._jornadas()[_momento(2, 0).date()]), "Corregida")

    def test_una_marca_anulada_no_cuenta(self):
        self._marca(_momento(2, 8), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 9), Marca.Tipo.SALIDA, anulada=True)
        self._marca(_momento(2, 16), Marca.Tipo.SALIDA)
        tramos = self._jornadas()[_momento(2, 0).date()]
        self.assertEqual(len(tramos), 1)
        self.assertEqual(tramos[0].horas, timedelta(hours=8))

    def test_un_mes_sin_marcas_tiene_todos_los_dias(self):
        jornadas = self._jornadas()
        self.assertEqual(len(jornadas), 30)
        self.assertTrue(all(observacion(tramos) == "Sin marcas" for tramos in jornadas.values()))

    def test_la_salida_del_dia_1_que_cierra_el_mes_anterior_no_aparece(self):
        self._marca(timezone.make_aware(_momento(1, 0).replace(tzinfo=None) - timedelta(hours=2)), Marca.Tipo.ENTRADA)
        self._marca(_momento(1, 5), Marca.Tipo.SALIDA)
        self.assertEqual(self._jornadas()[_momento(1, 0).date()], [])

    def test_las_horas_y_el_total_quedan_como_horas_en_el_archivo(self):
        self._turno(2, 8, 16)
        self._turno(3, 8, 12)
        hoja = load_workbook(BytesIO(generar_excel([(self.empleado, self.empleado.marcas.all())], 2026, 11)))["Juan Perez"]
        # Dia, Horario (US-14.14) y dos pares de entrada y salida.
        columna_horas = 3 + 2 * 2
        self.assertEqual(hoja.cell(row=5, column=columna_horas).value, "Horas del día")
        self.assertEqual(hoja.cell(row=FILA_DIA_1 + 1, column=3).value.strftime("%H:%M"), "08:00")
        self.assertEqual(hoja.cell(row=FILA_DIA_1 + 1, column=columna_horas).value, timedelta(hours=8))
        self.assertEqual(hoja.cell(row=FILA_DIA_1 + 30, column=1).value, "Total del mes")
        self.assertEqual(hoja.cell(row=FILA_DIA_1 + 30, column=columna_horas).value, timedelta(hours=12))

    def test_las_horas_del_dia_coinciden_con_las_horas_que_se_ven(self):
        # Los escaneos reales tienen segundos; la planilla muestra 08:00 y 16:00.
        self._marca(_momento(2, 8) + timedelta(seconds=50), Marca.Tipo.ENTRADA)
        self._marca(_momento(2, 16) + timedelta(seconds=5), Marca.Tipo.SALIDA)
        self.assertEqual(self._jornadas()[_momento(2, 0).date()][0].horas, timedelta(hours=8))

    def test_el_nombre_de_hoja_siempre_lo_acepta_excel(self):
        usados = set()
        for nombre in (
            "Ana Maria Fernanda Gonzalez D' Angelo",  # el corte cae en "' "
            "Juan Carlos Perez Gonzalez \U0001F600 Lopez",  # el emoji ocupa dos unidades
            "'Citado'",
            "''",
        ):
            with self.subTest(nombre=nombre):
                titulo = _titulo_hoja(nombre, usados)
                self.assertLessEqual(len(titulo.encode("utf-16-le")) // 2, 31)
                self.assertFalse(titulo.startswith("'") or titulo.endswith("'"))
                self.assertTrue(titulo)
        repetido = _titulo_hoja("Juan Carlos Perez Gonzalez \U0001F600 Lopez", usados)
        self.assertLessEqual(len(repetido.encode("utf-16-le")) // 2, 31)
        self.assertTrue(repetido.endswith("(2)"))

    def test_un_nombre_con_igual_no_se_vuelve_formula(self):
        self.empleado.nombre = "=HYPERLINK(\"http://x\")"
        contenido = generar_excel([(self.empleado, [])], 2026, 11)
        libro = load_workbook(BytesIO(contenido))
        hoja = libro.worksheets[1]
        self.assertEqual(hoja["A1"].data_type, "s")
        self.assertEqual(hoja["A1"].value, self.empleado.nombre)
        # En el Resumen tambien queda como texto.
        self.assertEqual(libro["Resumen"]["A6"].data_type, "s")
        self.assertEqual(libro["Resumen"]["A6"].value, self.empleado.nombre)


class PantallaDelExcelTests(ExcelTestCase):
    def _descargar(self, **params):
        return self.client.get(reverse("asistencia:excel"), {"descargar": "1", "mes": "2026-11", **params})

    def _hojas(self, response):
        # La primera es siempre el Resumen (US-14.11); despues, una por empleado.
        hojas = load_workbook(BytesIO(response.content)).sheetnames
        self.assertEqual(hojas[0], "Resumen")
        return hojas[1:]

    def test_una_hoja_por_empleado_del_alcance(self):
        self._turno(2, 8, 16)
        de_mapogo = self._empleado(nombre="De Mapogo", dni="31000111", sucursal=self.vivre)
        self._turno(2, 8, 16, empleado=de_mapogo)
        self._login(self.admin, [self.armadi])
        response = self._descargar()
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(self._hojas(response), ["Juan Perez"])

    def test_los_de_baja_entran_solo_si_tienen_marcas_ese_mes(self):
        sin_marcas = self._empleado(nombre="Ana Baja", dni="32000111")
        con_marcas = self._empleado(nombre="Beto Baja", dni="33000111")
        self._turno(3, 8, 16, empleado=con_marcas)
        dar_de_baja_empleado(actor=self.admin, empleado=sin_marcas)
        dar_de_baja_empleado(actor=self.admin, empleado=con_marcas)
        self._login(self.admin, [self.armadi])
        self.assertEqual(self._hojas(self._descargar()), ["Beto Baja", "Juan Perez"])

    def test_un_solo_empleado(self):
        self._empleado(nombre="Otro", dni="34000111")
        self._login(self.admin, [self.armadi])
        response = self._descargar(empleado=self.empleado.pk)
        self.assertEqual(self._hojas(response), ["Juan Perez"])
        self.assertIn("juan-perez", response["Content-Disposition"])

    def test_un_empleado_de_otra_empresa_no_se_puede_pedir(self):
        de_mapogo = self._empleado(nombre="De Mapogo", dni="31000111", sucursal=self.vivre)
        self._login(self.admin, [self.armadi])
        self.assertEqual(self._descargar(empleado=de_mapogo.pk).status_code, 400)

    def test_un_mes_invalido_no_descarga(self):
        self._login(self.admin, [self.armadi])
        for mes in ("2026-13", "2026-1²", "²026-11", "2026-11-01", "noviembre", ""):
            with self.subTest(mes=mes):
                self.assertEqual(self._descargar(mes=mes).status_code, 400)

    def test_un_dado_de_baja_cuya_unica_marca_cierra_el_mes_anterior_no_lleva_hoja(self):
        nocturno = self._empleado(nombre="Nocturno Baja", dni="36000111")
        self._marca(_momento(1, 0) - timedelta(hours=2), Marca.Tipo.ENTRADA, empleado=nocturno)
        self._marca(_momento(1, 6), Marca.Tipo.SALIDA, empleado=nocturno)
        dar_de_baja_empleado(actor=self.admin, empleado=nocturno)
        self._login(self.admin, [self.armadi])
        self.assertEqual(self._hojas(self._descargar()), ["Juan Perez"])
        # Pedido de a uno, sale igual: lo eligio quien descarga.
        self.assertEqual(self._hojas(self._descargar(empleado=nocturno.pk)), ["Nocturno Baja"])

    def test_quien_solo_ficha_no_entra(self):
        self._login(self.cajero, [self.armadi])
        self.assertEqual(self.client.get(reverse("asistencia:excel")).status_code, 403)
